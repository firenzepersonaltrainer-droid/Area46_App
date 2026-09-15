#!/usr/bin/env python3
"""
Riorganizzazione Strutturale e Visiva del Livello PRO in Cicli (A-B-A-B-A-B).
Area46 Training Lab - Landmine Lab PRO.

Logica:
- Le settimane sono sostituite dai Cicli (Ciclo 1, Ciclo 2).
- In ogni ciclo si alternano per 3 volte gli allenamenti A e B (A-B-A-B-A-B, 6 sessioni per ciclo).
- Alternanza cromatica delle righe:
    * Allenamento A: Grigio chiaro (#F2F2F2)
    * Allenamento B: Bianco (#FFFFFF)
- Bordo inferiore di separazione marcato sull'ultima riga di ogni ciclo.
- Applicato sia su Google Fogli (entrambi i fogli) che sui file Excel locali e demo-data.json.
"""

import os
import json
import gspread
import openpyxl
from openpyxl.styles import PatternFill, Border, Side, Font, Alignment

CREDENTIALS_FILE = "credenziali.json"
TARGET_SHEET_ID = "1AQHkYjXhI6WhRPORJ4DntiGZQd8ob54wJxkwP_LymwY" # PRO database esercizi
LAB_SHEET_ID = "1SNy8IJk1E20BcAdar0F1TqqnpRIDkjlkNn6P2ljKPNY"    # database esercizi Lab
LOCAL_LAB_EXCEL = "database esercizi Lab.xlsx"
LOCAL_PRO_EXCEL = "PRO database esercizi.xlsx"
DEMO_DATA_JSON = "demo-data.json"

COLUMNS = [
    "livello",
    "Settimana",
    "giorno",
    "Sequenza",
    "ID esercizio",
    "nome esercizio",
    "parametri",
    "recupero",
    "minutaggio blocco",
    "Note Tecniche",
    "video temp",
    "data_pubblicazione"
]

# --- DEFINIZIONE WORKOUT CARDINE ---

# CICLO 1: Allenamento A (9 esercizi)
C1_A = [
    ("1a", 51, "High hand switch", "3x30", "", "jump set"),
    ("1b", 54, "Landmine press Cross Step", "3x8+8", "", "jump set"),
    ("2",  56, "Landmine coiled Deadlift", "3x6+6", 'RBS 60"', "rpe 8"),
    ("3",  43, "Pounce squat", "4x8", '90"', "rpe 8"),
    ("4a", 82, "Landmine Biker's Squat", "3x10", "", "jump set"),
    ("4b", 81, "Meadow Row", "3x10+10", '60"', "jump set"),
    ("5a", 69, "Split Snatch", '5x30"', "", "circuito"),
    ("5b", 170, "Coiled Kettlebell Swing", '5x30"', "", "circuito"),
    ("5c", 100, "Speed skater touch", '5x30"', '45"', "circuito")
]

# CICLO 1: Allenamento B (10 esercizi)
C1_B = [
    ("1a", 46, "Split Switch Screwdriver", "3x30", "", "jump set"),
    ("1b", 64, "High Pull ISO", '3x15"+15"', "", "jump set"),
    ("2a", 59, "Split Clean and Jerk", "3x6+6 (rpe 8)", 'RBS 45"', "jump set"),
    ("2b", 45, "Landmine rotational Lunges", "3x8 (RPE8)", '45"', "jump set"),
    ("3a", 230, "kettlebell clean and press", "3x6+6", 'RBS 45"', "jump set"),
    ("3b", 44, "Coiled Reverse Lunge", "3x6+6", 'RBS 45"', "jump set"),
    ("4a", 205, "Kick-through", '4x30"', '60"', "jump set"),
    ("4b", 206, "Side kick through", '4x30"', '60"', "jump set"),
    ("5a", 63, "Landmine Clean and jerk NO DOWN", "x 8+8", "", "AMRAP 10'"),
    ("5b", 203, "Plyometric lunges", "x 10", "", "AMRAP 10'")
]

# CICLO 2: Allenamento A (8 esercizi - Spinta orizzontale pesante & Trazione)
C2_A = [
    ("1a", 98, "Avenger Blade passing", "3x24", "", "jump set"),
    ("1b", 49, "Step/Step back press", "3x8+8", '45"', "jump set"),
    ("2",  79, "Bench press", "4x8", '90"', "RPE 8"),
    ("3a", 81, "Meadow Row", "4x8+8", '60"', "jump set | RPE 8"),
    ("3b", 82, "Landmine Biker's Squat", "4x8+8", 'RBS 60"', "jump set | RPE 8"),
    ("4a", 172, "Kettlebell swing to clean", "4x8+8", "", "circuito (4 round)"),
    ("4b", 124, "Slamball Skater jump", "4x10", "", "circuito (4 round)"),
    ("4c", 1,   "skierg regular", '4x40"', '60"', "circuito (4 round)")
]

# CICLO 2: Allenamento B (9 esercizi - Trazione manubri & Spinta orizzontale & AMRAP 15')
C2_B = [
    ("1a", 55, "Hang Position ISO", '3x20"+20"', "", "jump set"),
    ("1b", 116, "High pull bounce", "3x10+10", '45"', "jump set"),
    ("2a", 53, "Split C&J Touch & Go", "4x6+6", 'RBS 45"', "jump set | RPE 8"),
    ("2b", 76, "Push the wall row", "4x8+8", 'RBS 45"', "jump set | RPE 8"),
    ("3a", 101, "Sprinter Push up", "4x10", '45"', "jump set | RPE 8"),
    ("3b", 44, "Coiled Reverse Lunge", "4x8+8", 'RBS 45"', "jump set | RPE 8"),
    ("4a", 73, "Lateral landmine clean", "x 6+6", "", "AMRAP 15'"),
    ("4b", 206, "Side kick through", "x 8+8", "", "AMRAP 15'"),
    ("4c", 187, "Sprawl to Broad Jump", "x 8", "", "AMRAP 15'")
]


def generate_all_rows():
    """
    Genera le 108 righe complessive per i due Cicli.
    Ritorna:
      all_rows: lista di liste con i valori delle 12 colonne
      meta_sessions: lista di dizionari con info su ogni riga (per la colorazione e i bordi)
    """
    all_rows = []
    meta = [] # per ogni riga: {'is_a': bool, 'is_cycle_end': bool}

    # CICLO 1: A-B-A-B-A-B (sessioni 1-6)
    c1_plan = [
        (1, "A", C1_A),
        (2, "B", C1_B),
        (3, "A", C1_A),
        (4, "B", C1_B),
        (5, "A", C1_A),
        (6, "B", C1_B)
    ]

    for sess_num, sess_type, exercises in c1_plan:
        giorno_str = f"{sess_num} allenamento {sess_type}"
        is_a = (sess_type == "A")
        for i, ex in enumerate(exercises):
            is_last_in_cycle = (sess_num == 6 and i == len(exercises) - 1)
            row = [
                "PRO",
                "Ciclo 1",
                giorno_str,
                ex[0], # seq
                ex[1], # id
                ex[2], # nome
                ex[3], # param
                ex[4], # rec
                "",    # minutaggio
                ex[5], # note
                "",    # video
                ""     # data_pub
            ]
            all_rows.append(row)
            meta.append({"is_a": is_a, "is_cycle_end": is_last_in_cycle})

    # CICLO 2: A-B-A-B-A-B (sessioni 7-12)
    c2_plan = [
        (7,  "A", C2_A),
        (8,  "B", C2_B),
        (9,  "A", C2_A),
        (10, "B", C2_B),
        (11, "A", C2_A),
        (12, "B", C2_B)
    ]

    for sess_num, sess_type, exercises in c2_plan:
        giorno_str = f"{sess_num} allenamento {sess_type}"
        is_a = (sess_type == "A")
        for i, ex in enumerate(exercises):
            is_last_in_cycle = (sess_num == 12 and i == len(exercises) - 1)
            row = [
                "PRO",
                "Ciclo 2",
                giorno_str,
                ex[0], # seq
                ex[1], # id
                ex[2], # nome
                ex[3], # param
                ex[4], # rec
                "",    # minutaggio
                ex[5], # note
                "",    # video
                ""     # data_pub
            ]
            all_rows.append(row)
            meta.append({"is_a": is_a, "is_cycle_end": is_last_in_cycle})

    return all_rows, meta


def apply_google_sheet_styling(spreadsheet, worksheet, meta):
    """
    Applica alternanza colori (A grigio chiaro, B bianco)
    e bordo inferiore marcato di separazione tra i cicli.
    """
    sheet_id = worksheet.id
    requests = []

    # 1. Header styling (sfondo scuro o leggero, testo in grassetto)
    requests.append({
        "repeatCell": {
            "range": {
                "sheetId": sheet_id,
                "startRowIndex": 0,
                "endRowIndex": 1,
                "startColumnIndex": 0,
                "endColumnIndex": len(COLUMNS)
            },
            "cell": {
                "userEnteredFormat": {
                    "textFormat": {"bold": True},
                    "backgroundColor": {"red": 0.88, "green": 0.88, "blue": 0.88}
                }
            },
            "fields": "userEnteredFormat(textFormat,backgroundColor)"
        }
    })

    # 2. Colorazione alternata delle righe (Allenamento A grigio chiaro, B bianco)
    # Raggruppiamo blocchi contigui per minimizzare le richieste API
    idx = 0
    while idx < len(meta):
        is_a = meta[idx]["is_a"]
        start_row = idx + 1 # +1 per l'header
        while idx < len(meta) and meta[idx]["is_a"] == is_a:
            idx += 1
        end_row = idx + 1

        bg_color = {"red": 0.95, "green": 0.95, "blue": 0.95} if is_a else {"red": 1.0, "green": 1.0, "blue": 1.0}
        requests.append({
            "repeatCell": {
                "range": {
                    "sheetId": sheet_id,
                    "startRowIndex": start_row,
                    "endRowIndex": end_row,
                    "startColumnIndex": 0,
                    "endColumnIndex": len(COLUMNS)
                },
                "cell": {
                    "userEnteredFormat": {
                        "backgroundColor": bg_color
                    }
                },
                "fields": "userEnteredFormat.backgroundColor"
            }
        })

    # 3. Bordi inferiori marcati alla fine di ogni Ciclo
    for idx, m in enumerate(meta):
        if m["is_cycle_end"]:
            row_idx = idx + 1 # +1 per l'header (0-indexed per startRowIndex, endRowIndex = row_idx + 1)
            requests.append({
                "updateBorders": {
                    "range": {
                        "sheetId": sheet_id,
                        "startRowIndex": row_idx,
                        "endRowIndex": row_idx + 1,
                        "startColumnIndex": 0,
                        "endColumnIndex": len(COLUMNS)
                    },
                    "bottom": {
                        "style": "SOLID_MEDIUM",
                        "width": 2,
                        "color": {"red": 0.0, "green": 0.0, "blue": 0.0}
                    }
                }
            })

    # Esegui batchUpdate
    body = {"requests": requests}
    spreadsheet.batch_update(body)
    print(f"-> Styling applicato con successo sul foglio '{worksheet.title}' (ID {sheet_id})!")


def update_excel_file(filepath, sheet_name, all_rows, meta):
    """Aggiorna e formatta graficamente il file Excel locale con openpyxl."""
    if not os.path.exists(filepath):
        print(f"File {filepath} non trovato, salto.")
        return

    wb = openpyxl.load_workbook(filepath)
    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(title=sheet_name)
    else:
        ws = wb[sheet_name]

    # Svuota righe esistenti
    for r in range(1, max(ws.max_row + 1, 150)):
        for c in range(1, 15):
            cell = ws.cell(row=r, column=c)
            cell.value = None
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border()

    # Scrivi Header
    header_fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
    header_font = Font(bold=True)
    for col_idx, col_name in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font

    # Stili
    fill_a = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    fill_b = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    border_bottom_thick = Border(bottom=Side(style="medium", color="000000"))

    # Scrivi Dati e Stili
    for row_idx, (r_vals, m) in enumerate(zip(all_rows, meta), 2):
        row_fill = fill_a if m["is_a"] else fill_b
        is_end = m["is_cycle_end"]
        for col_idx, val in enumerate(r_vals, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = row_fill
            if is_end:
                cell.border = border_bottom_thick

    wb.save(filepath)
    print(f"-> File Excel locale '{filepath}' (foglio '{sheet_name}') aggiornato e formattato!")


def main():
    print("=" * 75)
    print("Area46 Training Lab - Riorganizzazione Livello PRO in Cicli (A-B-A-B-A-B)")
    print("=" * 75)

    all_rows, meta = generate_all_rows()
    print(f"\nGenerate {len(all_rows)} righe complessive:")
    print("  - Ciclo 1 (1-6): 57 righe (3x Allenamento A [9 es.] e 3x Allenamento B [10 es.])")
    print("  - Ciclo 2 (7-12): 51 righe (3x Allenamento A [8 es.] e 3x Allenamento B [9 es.])")

    # 1. Connessione Google API
    print("\n[Fase 1] Connessione a Google Sheets...")
    gc = gspread.service_account(filename=CREDENTIALS_FILE)

    # 2. Aggiornamento Google Sheet TARGET (PRO database esercizi)
    print(f"\n[Fase 2] Aggiornamento Google Sheet 'PRO database esercizi' ({TARGET_SHEET_ID})...")
    sh_target = gc.open_by_key(TARGET_SHEET_ID)
    
    # Assicuriamoci che esista la scheda 'PRO' (o rinominiamo Allenamenti PRO se presente)
    try:
        ws_pro_target = sh_target.worksheet("PRO")
    except gspread.WorksheetNotFound:
        try:
            ws_pro_target = sh_target.worksheet("Allenamenti PRO")
            ws_pro_target.update_title("PRO")
            print("Scheda 'Allenamenti PRO' rinominata in 'PRO'.")
        except gspread.WorksheetNotFound:
            ws_pro_target = sh_target.add_worksheet(title="PRO", rows=150, cols=15)
            print("Creata nuova scheda 'PRO'.")

    # Pulizia completa e scrittura
    ws_pro_target.batch_clear(["A1:L200"])
    data_to_write = [COLUMNS] + all_rows
    ws_pro_target.update(range_name=f"A1:L{len(data_to_write)}", values=data_to_write)
    print(f"-> Scritte {len(data_to_write)} righe (header + dati) nella scheda 'PRO'.")
    
    # Applicazione formattazione visiva
    apply_google_sheet_styling(sh_target, ws_pro_target, meta)

    # 3. Aggiornamento Google Sheet MASTER (database esercizi Lab)
    print(f"\n[Fase 3] Aggiornamento Google Sheet 'database esercizi Lab' ({LAB_SHEET_ID})...")
    sh_lab = gc.open_by_key(LAB_SHEET_ID)
    ws_pro_lab = sh_lab.worksheet("Pro")
    
    ws_pro_lab.batch_clear(["A1:L200"])
    ws_pro_lab.update(range_name=f"A1:L{len(data_to_write)}", values=data_to_write)
    print(f"-> Scritte {len(data_to_write)} righe (header + dati) nella scheda 'Pro'.")
    
    # Applicazione formattazione visiva
    apply_google_sheet_styling(sh_lab, ws_pro_lab, meta)

    # 4. Aggiornamento file Excel locali con openpyxl
    print("\n[Fase 4] Aggiornamento e formattazione dei file Excel locali...")
    update_excel_file(LOCAL_LAB_EXCEL, "Pro", all_rows, meta)
    update_excel_file(LOCAL_PRO_EXCEL, "PRO", all_rows, meta)

    # 5. Aggiornamento demo-data.json
    print("\n[Fase 5] Sincronizzazione demo-data.json...")
    if os.path.exists(DEMO_DATA_JSON):
        with open(DEMO_DATA_JSON, "r") as f:
            demo_data = json.load(f)

        # Rimuoviamo vecchi record PRO
        demo_data["allenamenti"] = [a for a in demo_data["allenamenti"] if a.get("livello") != "PRO"]

        existing_ids = [a.get("id", 0) for a in demo_data["allenamenti"]]
        max_id = max(existing_ids) if existing_ids else 500

        for r in all_rows:
            max_id += 1
            demo_data["allenamenti"].append({
                "id": max_id,
                "livello": r[0],
                "settimana": r[1],
                "giorno": r[2],
                "sequenza": r[3],
                "id_esercizio": f"{int(r[4]):03d}",
                "nome_esercizio": r[5],
                "parametri": r[6],
                "recupero": r[7],
                "minutaggio_blocco": "",
                "note_tecniche": r[9],
                "link_video": "",
                "data_pubblicazione": None
            })

        with open(DEMO_DATA_JSON, "w") as f:
            json.dump(demo_data, f, indent=2)
        print(f"-> File '{DEMO_DATA_JSON}' sincronizzato con {len(all_rows)} allenamenti PRO.")

    print("\n" + "=" * 75)
    print("RIORGANIZZAZIONE IN CICLI E FORMATTAZIONE GRAFICA COMPLETATE CON SUCCESSO!")
    print("=" * 75)


if __name__ == "__main__":
    main()
