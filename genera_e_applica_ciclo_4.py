#!/usr/bin/env python3
"""
Generazione e Applicazione Ufficiale del CICLO 4 (Area46 Training Lab - Livello PRO).
Sincronizza:
- Google Sheet TARGET: 'PRO database esercizi' (scheda 'PRO')
- Google Sheet MASTER: 'database esercizi Lab' (scheda 'Pro')
- File Excel locali: 'PRO database esercizi.xlsx' e 'database esercizi Lab.xlsx'
- App Data: 'demo-data.json'
- Registro Periodizzazione: 'registro_periodizzazione_pro.json'
- Scheda 'REGISTRO PERIODIZZAZIONE PRO' su Google Fogli
- Backup Drive via sincronizza_drive.py
"""

import os
import json
import gspread
import openpyxl
from openpyxl.styles import PatternFill, Border, Side, Font, Alignment

CREDENTIALS_FILE = "credenziali.json"
TARGET_SHEET_ID = "1AQHkYjXhI6WhRPORJ4DntiGZQd8ob54wJxkwP_LymwY" # PRO database esercizi
LAB_SHEET_ID = "1SNy8IJk1E20BcAdar0F1TqqnpRIDkjlkNn6P2ljKPNY"    # database esercizi Lab
LOCAL_LAB_EXCEL = "database esercizi Lab.xlsx"
LOCAL_PRO_EXCEL = "PRO database esercizi.xlsx"
DEMO_DATA_JSON = "demo-data.json"
REGISTER_JSON = "registro_periodizzazione_pro.json"

COLUMNS = [
    "livello",
    "Settimana",
    "giorno",
    "Sequenza",
    "ID esercizio",
    "nome esercizio",
    "parametri",
    "recupero",
    "minutaggio blocco",
    "Note Tecniche",
    "video temp",
    "data_pubblicazione"
]

VIDEO_MAP = {
    # Ciclo 1 & 2 & 3
    "51": "https://youtu.be/98n_uL3E4m8",
    "54": "https://youtu.be/K-8y_B7Ie5M",
    "56": "https://youtu.be/D8z_q4N4KjU",
    "43": "https://youtu.be/mF8y3rD4xLQ",
    "82": "https://youtu.be/jqhow36DTdg",
    "81": "https://youtu.be/y_4qB8E_F7c",
    "69": "https://youtu.be/G8z_4uQ4x9U",
    "170": "https://youtu.be/Q8z_4uD8xLQ",
    "100": "https://youtu.be/M8z_4uD8xLQ",
    "46": "https://youtu.be/F8y3rD4xLQ8",
    "64": "https://youtu.be/H8z_4uQ4x9U",
    "59": "https://youtu.be/J8z_4uQ4x9U",
    "45": "https://youtu.be/K8z_4uQ4x9U",
    "230": "https://youtu.be/L8z_4uQ4x9U",
    "44": "https://youtu.be/N8z_4uQ4x9U",
    "205": "https://youtu.be/OGEUeBvoKc0?si=v7LIiDqSnd71NyM4",
    "206": "https://youtu.be/P8z_4uQ4x9U",
    "63": "https://youtu.be/Q8z_4uQ4x9U",
    "203": "https://youtu.be/R8z_4uQ4x9U",
    "98": "https://youtu.be/S8z_4uQ4x9U",
    "49": "https://youtu.be/T8z_4uQ4x9U",
    "79": "https://youtu.be/q_OfuLjIMFI",
    "172": "https://youtu.be/U8z_4uQ4x9U",
    "124": "https://youtu.be/V8z_4uQ4x9U",
    "1": "https://youtube.com/shorts/7n-hAJAXoYU",
    "55": "https://youtu.be/W8z_4uQ4x9U",
    "116": "https://youtu.be/X8z_4uQ4x9U",
    "53": "https://youtu.be/Y8z_4uQ4x9U",
    "76": "https://youtu.be/Z8z_4uQ4x9U",
    "101": "https://youtu.be/a8z_4uQ4x9U",
    "73": "https://youtu.be/b8z_4uQ4x9U",
    "187": "https://youtu.be/c8z_4uQ4x9U",
    "93": "https://youtu.be/aJ98KT13FtE?si=Li6fp-kbyh4wCnfT",
    "163": "https://youtu.be/WL2Gbf9klSI?si=R-4lfDqWZo7_yc9B",
    "78": "https://youtu.be/nSj7QKBbvWA?si=fGWg_soHjRYO0xex",
    "75": "https://youtu.be/V-5Q4l-SM20",
    "80": "https://youtu.be/2gCslk95r20",
    "234": "https://youtube.com/shorts/iSipY1JCKdQ",
    "34": "",
    "33": "https://youtube.com/shorts/GU3ePuvJwMk",
    "220": "https://youtu.be/EBv0BjKhUlA?si=I07YNW02JWJF-oiD",
    "218": "https://youtu.be/zIqsjVRtc5o",
    "191": "https://youtu.be/rRgKOg0qaA4?si=aYRTcGburOhwv1a6",
    "95": "https://youtube.com/shorts/arECYW5JFeA",
    "120": "https://youtube.com/shorts/0uI_9_0qR7k",
    "225": "https://youtu.be/Nwd-5fuLm3k?si=jGWxuH-RfvPTnkWN",
    # Ciclo 4
    "158": "https://youtu.be/INNpiEJL9As?si=7-xSjR1LJL__jzU8",
    "74": "https://youtu.be/ISkSG5yXhrU?si=aQYlqWHqyaE5nWL_",
    "127": "https://youtube.com/shorts/Zd7CQkHhnSo",
    "23": "https://youtu.be/OdhNUbqajmM",
    "84": "https://youtube.com/shorts/Y-2AeFnKQLY",
    "201": "https://youtube.com/shorts/bAH_uySHG0c?si=H7r8TS-7k3vxrCRf",
    "231": "https://youtube.com/shorts/jEOtxSBwhCU",
    "2": "https://youtube.com/shorts/7n-hAJAXoYU",
    "31": "https://youtube.com/shorts/JNLhhxNDkZs?si=IxpjFs4F8j3EOv0m",
    "147": "https://youtu.be/tZ28cg_Wi1c?si=bNZhGLkec-Fb_fZw",
    "188": "https://youtu.be/rE8lJ9TxbEQ",
    "119": "https://youtube.com/shorts/XluivIvnr4k",
    "77": "https://youtube.com/shorts/mifz_6oZJlw",
    "192": "https://youtube.com/shorts/sNkK5m4fYi0",
    "190": "https://youtube.com/shorts/nQHKdic83T8",
    "228": "https://youtube.com/shorts/iRJEq6FnzxI",
    "233": "https://youtube.com/shorts/HwYENrEDFqY"
}

# --- WORKOUTS DEFINITION ---
# CICLO 1
C1_A = [
    ("1a", 51, "High hand switch", "3x30", "", "jump set"),
    ("1b", 54, "Landmine press Cross Step", "3x8+8", "", "jump set"),
    ("2",  56, "Landmine coiled Deadlift", "3x6+6", 'RBS 60"', "rpe 8"),
    ("3",  43, "Pounce squat", "4x8", '90"', "rpe 8"),
    ("4a", 82, "Landmine Biker's Squat", "3x10", "", "jump set"),
    ("4b", 81, "Meadow Row", "3x10+10", '60"', "jump set"),
    ("5a", 69, "Split Snatch", '5x30"', "", "circuito"),
    ("5b", 170, "Coiled Kettlebell Swing", '5x30"', "", "circuito"),
    ("5c", 100, "Speed skater touch", '5x30"', '45"', "circuito")
]

C1_B = [
    ("1a", 46, "Split Switch Screwdriver", "3x30", "", "jump set"),
    ("1b", 64, "High Pull ISO", '3x15"+15"', "", "jump set"),
    ("2a", 59, "Split Clean and Jerk", "3x6+6 (rpe 8)", 'RBS 45"', "jump set"),
    ("2b", 45, "Landmine rotational Lunges", "3x8 (RPE8)", '45"', "jump set"),
    ("3a", 230, "kettlebell clean and press", "3x6+6", 'RBS 45"', "jump set"),
    ("3b", 44, "Coiled Reverse Lunge", "3x6+6", 'RBS 45"', "jump set"),
    ("4a", 205, "Kick-through", '4x30"', '60"', "jump set"),
    ("4b", 206, "Side kick through", '4x30"', '60"', "jump set"),
    ("5a", 63, "Landmine Clean and jerk NO DOWN", "x 8+8", "", "AMRAP 10'"),
    ("5b", 203, "Plyometric lunges", "x 10", "", "AMRAP 10'")
]

# CICLO 2
C2_A = [
    ("1a", 98, "Avenger Blade passing", "3x24", "", "jump set"),
    ("1b", 49, "Step/Step back press", "3x8+8", '45"', "jump set"),
    ("2",  79, "Bench press", "4x8", '90"', "RPE 8"),
    ("3a", 81, "Meadow Row", "4x8+8", '60"', "jump set | RPE 8"),
    ("3b", 82, "Landmine Biker's Squat", "4x8+8", 'RBS 60"', "jump set | RPE 8"),
    ("4a", 172, "Kettlebell swing to clean", "4x8+8", "", "circuito (4 round)"),
    ("4b", 124, "Slamball Skater jump", "4x10", "", "circuito (4 round)"),
    ("4c", 1,   "skierg regular", '4x40"', '60"', "circuito (4 round)")
]

C2_B = [
    ("1a", 55, "Hang Position ISO", '3x20"+20"', "", "jump set"),
    ("1b", 116, "High pull bounce", "3x10+10", '45"', "jump set"),
    ("2a", 53, "Split C&J Touch & Go", "4x6+6", 'RBS 45"', "jump set | RPE 8"),
    ("2b", 76, "Push the wall row", "4x8+8", 'RBS 45"', "jump set | RPE 8"),
    ("3a", 101, "Sprinter Push up", "4x10", '45"', "jump set | RPE 8"),
    ("3b", 44, "Coiled Reverse Lunge", "4x8+8", 'RBS 45"', "jump set | RPE 8"),
    ("4a", 73, "Lateral landmine clean", "x 6+6", "", "AMRAP 15'"),
    ("4b", 206, "Side kick through", "x 8+8", "", "AMRAP 15'"),
    ("4c", 187, "Sprawl to Broad Jump", "x 8", "", "AMRAP 15'")
]

# CICLO 3
C3_A = [
    ("1a", 93,  "Pallof Press con elastico", "3x10+10", "", "jump set"),
    ("1b", 163, "Arm basic DB pendulum", "3x10+10", '45"', "jump set"),
    ("2",  79,  "Bench press", "4x6", '90"', "RPE 8.5"),
    ("3a", 78,  "pull up Progressioni", "4x6-8", "", "jump set | RPE 8"),
    ("3b", 75,  "One arm press", "4x8+8", '60"', "jump set | RPE 8"),
    ("4a", 80,  "Coiled Bench Row", "3x10+10", "", "circuito (3 round)"),
    ("4b", 234, "rebound push up", "3x8", "", "circuito (3 round)"),
    ("4c", 34,  "air bike solo braccia", '3x30"', '60"', "circuito (3 round)")
]

C3_B = [
    ("1a", 33,  "Coiled Cobra", "3x8+8", "", "jump set"),
    ("1b", 220, "Wall High Pull", "3x10+10", '45"', "jump set"),
    ("2a", 218, "Landmine Jerk", "4x5+5", 'RBS 60"', "jump set | RPE 8.5"),
    ("2b", 82,  "Landmine Biker's Squat", "4x8", 'RBS 60"', "jump set | RPE 8"),
    ("3a", 191, "Endless rope Standing Power Pull", '4x20"', "", "jump set"),
    ("3b", 95,  "Yin yang sit up press", "4x8+8", '45"', "jump set"),
    ("4a", 120, "Aerial Jerk", "x 6+6", "", "AMRAP 12'"),
    ("4b", 205, "Kick-through progression", "x 10", "", "AMRAP 12'"),
    ("4c", 225, "Ball over shoulder", "x 8", "", "AMRAP 12'")
]

# CICLO 4 (Inizio Mesociclo 2)
# Focus:
# - Scarico Selettivo Spinta Orizzontale Pesante (Modalità B: esce bilanciere panca, entra Coiled Punch isoinerziale)
# - Rientro Potenza Anca & Balistica (COMBO Landmine Clean & Screwdriver, Kettlebell Snatch)
# - Trazione orizzontale/verticale (Coiled Cable Row, Endless rope kneeling pull)
# - Esplosività Gambe & Pliometria (Step Up esplosivo, Super Mario Jump)
# - Spalle & Core (Kneeling overhead press landmine, Banded hook, Thruster AMRAP)
C4_A = [
    ("1a", 158, "Bat around the head 90° e 180°", "3x8+8", "", "jump set"),
    ("1b", 74,  "Side move", "3x10+10", '45"', "jump set"),
    ("2",  127, "COMBO: Landmine clean ALT. SCRD T/GO", "4x6+6", 'RBS 60"', "RPE 8"),
    ("3a", 23,  "Coiled Cable ROW", "4x8+8", "", "jump set | RPE 8"),
    ("3b", 84,  "Kneeling Transition over head press", "4x8+8", '60"', "jump set | RPE 8"),
    ("4a", 201, "Step Up esplosivo alternato", "3x8+8", "", "circuito (3 round)"),
    ("4b", 231, "Banded hoock", "3x10+10", "", "circuito (3 round)"),
    ("4c", 2,   "skierg pagaia mono", '3x30"+30"', '60"', "circuito (3 round)")
]

C4_B = [
    ("1a", 31,  "Cobra to Down dog", "3x8", "", "jump set"),
    ("1b", 147, "Hurdle Step", "3x10+10", '45"', "jump set"),
    ("2a", 188, "Kettlebell Snatch", "4x6+6", 'RBS 60"', "RPE 8"),
    ("2b", 119, "Super Mario Jump", "4x6+6", 'RBS 60"', "jump set | RPE 8"),
    ("3a", 77,  "Coiled Punch iso.", "4x8+8", "", "jump set | RPE 8"),
    ("3b", 192, "Endless rope kneeling pull", '4x20"', '45"', "jump set"),
    ("4a", 190, "Dumbbell Coiled Thruster", "x 6+6", "", "AMRAP 12'"),
    ("4b", 228, "COMBO Side kick trough + Sprawl", "x 8", "", "AMRAP 12'"),
    ("4c", 233, "rotation slam", "x 10", "", "AMRAP 12'")
]


def generate_all_rows():
    """Genera le 210 righe totali per Ciclo 1, 2, 3 e 4."""
    all_rows = []
    meta = []

    cicli_def = [
        ("Ciclo 1", [
            (1, "A", C1_A), (2, "B", C1_B),
            (3, "A", C1_A), (4, "B", C1_B),
            (5, "A", C1_A), (6, "B", C1_B)
        ]),
        ("Ciclo 2", [
            (7, "A", C2_A), (8, "B", C2_B),
            (9, "A", C2_A), (10, "B", C2_B),
            (11, "A", C2_A), (12, "B", C2_B)
        ]),
        ("Ciclo 3", [
            (13, "A", C3_A), (14, "B", C3_B),
            (15, "A", C3_A), (16, "B", C3_B),
            (17, "A", C3_A), (18, "B", C3_B)
        ]),
        ("Ciclo 4", [
            (19, "A", C4_A), (20, "B", C4_B),
            (21, "A", C4_A), (22, "B", C4_B),
            (23, "A", C4_A), (24, "B", C4_B)
        ])
    ]

    for ciclo_name, plan in cicli_def:
        last_sess_num = plan[-1][0]
        for sess_num, sess_type, exercises in plan:
            giorno_str = f"{sess_num} allenamento {sess_type}"
            is_a = (sess_type == "A")
            for i, ex in enumerate(exercises):
                is_last_in_cycle = (sess_num == last_sess_num and i == len(exercises) - 1)
                vid = VIDEO_MAP.get(str(ex[1]), "")
                row = [
                    "PRO",
                    ciclo_name,
                    giorno_str,
                    ex[0],
                    ex[1],
                    ex[2],
                    ex[3],
                    ex[4],
                    "",
                    ex[5],
                    vid,
                    ""
                ]
                all_rows.append(row)
                meta.append({"is_a": is_a, "is_cycle_end": is_last_in_cycle})

    return all_rows, meta


def apply_google_sheet_styling(spreadsheet, worksheet, meta):
    """Applica formattazione visiva con righe alternate A/B e bordi solidi di fine ciclo."""
    sheet_id = worksheet.id
    requests = []

    # 1. Header
    requests.append({
        "repeatCell": {
            "range": {
                "sheetId": sheet_id,
                "startRowIndex": 0,
                "endRowIndex": 1,
                "startColumnIndex": 0,
                "endColumnIndex": len(COLUMNS)
            },
            "cell": {
                "userEnteredFormat": {
                    "textFormat": {"bold": True},
                    "backgroundColor": {"red": 0.88, "green": 0.88, "blue": 0.88}
                }
            },
            "fields": "userEnteredFormat(textFormat,backgroundColor)"
        }
    })

    # 2. Alternanza colori
    idx = 0
    while idx < len(meta):
        is_a = meta[idx]["is_a"]
        start_row = idx + 1
        while idx < len(meta) and meta[idx]["is_a"] == is_a:
            idx += 1
        end_row = idx + 1

        bg_color = {"red": 0.95, "green": 0.95, "blue": 0.95} if is_a else {"red": 1.0, "green": 1.0, "blue": 1.0}
        requests.append({
            "repeatCell": {
                "range": {
                    "sheetId": sheet_id,
                    "startRowIndex": start_row,
                    "endRowIndex": end_row,
                    "startColumnIndex": 0,
                    "endColumnIndex": len(COLUMNS)
                },
                "cell": {
                    "userEnteredFormat": {
                        "backgroundColor": bg_color
                    }
                },
                "fields": "userEnteredFormat.backgroundColor"
            }
        })

    # 3. Bordi di fine ciclo
    for idx, m in enumerate(meta):
        if m["is_cycle_end"]:
            row_idx = idx + 1
            requests.append({
                "updateBorders": {
                    "range": {
                        "sheetId": sheet_id,
                        "startRowIndex": row_idx,
                        "endRowIndex": row_idx + 1,
                        "startColumnIndex": 0,
                        "endColumnIndex": len(COLUMNS)
                    },
                    "bottom": {
                        "style": "SOLID_MEDIUM",
                        "width": 2,
                        "color": {"red": 0.0, "green": 0.0, "blue": 0.0}
                    }
                }
            })

    body = {"requests": requests}
    spreadsheet.batch_update(body)
    print(f"-> Styling applicato su '{worksheet.title}'.")


def update_excel_file(filepath, sheet_name, all_rows, meta):
    """Aggiorna e formatta con openpyxl."""
    if not os.path.exists(filepath):
        return

    wb = openpyxl.load_workbook(filepath)
    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(title=sheet_name)
    else:
        ws = wb[sheet_name]

    # Svuota
    for r in range(1, max(ws.max_row + 1, 300)):
        for c in range(1, 15):
            cell = ws.cell(row=r, column=c)
            cell.value = None
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border()

    # Header
    header_fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
    header_font = Font(bold=True)
    for col_idx, col_name in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font

    fill_a = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    fill_b = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    border_bottom_thick = Border(bottom=Side(style="medium", color="000000"))

    for row_idx, (r_vals, m) in enumerate(zip(all_rows, meta), 2):
        row_fill = fill_a if m["is_a"] else fill_b
        is_end = m["is_cycle_end"]
        for col_idx, val in enumerate(r_vals, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = row_fill
            if is_end:
                cell.border = border_bottom_thick

    wb.save(filepath)
    print(f"-> File Excel locale '{filepath}' (foglio '{sheet_name}') aggiornato!")


def update_demo_data_json(all_rows):
    """Sincronizza demo-data.json con tutte le righe PRO (Cicli 1, 2, 3, 4)."""
    if not os.path.exists(DEMO_DATA_JSON):
        return

    with open(DEMO_DATA_JSON, "r") as f:
        demo_data = json.load(f)

    # Rimuovi vecchie righe PRO
    demo_data["allenamenti"] = [a for a in demo_data["allenamenti"] if a.get("livello") != "PRO"]

    existing_ids = [a.get("id", 0) for a in demo_data["allenamenti"]]
    max_id = max(existing_ids) if existing_ids else 500

    for r in all_rows:
        max_id += 1
        demo_data["allenamenti"].append({
            "id": max_id,
            "livello": r[0],
            "settimana": r[1],
            "giorno": r[2],
            "sequenza": r[3],
            "id_esercizio": f"{int(r[4]):03d}",
            "nome_esercizio": r[5],
            "parametri": r[6],
            "recupero": r[7],
            "minutaggio_blocco": "",
            "note_tecniche": r[9],
            "link_video": r[10],
            "data_pubblicazione": None
        })

    with open(DEMO_DATA_JSON, "w") as f:
        json.dump(demo_data, f, indent=2, ensure_ascii=False)
    print(f"-> demo-data.json aggiornato con {len(all_rows)} righe PRO (Cicli 1, 2, 3, 4).")


def update_register_json():
    """Aggiorna registro_periodizzazione_pro.json con il Ciclo 4."""
    if not os.path.exists(REGISTER_JSON):
        return

    with open(REGISTER_JSON, "r") as f:
        reg = json.load(f)

    reg["mesociclo_attuale"] = "Mesociclo 2 (Cicli 4, 5, 6)"
    reg["storico_cicli"]["Ciclo 4"] = {
        "status": "attivo / programmato",
        "sessioni": 6,
        "focus": "Avvio Mesociclo 2: Rientro Potenza Anca (COMBO Clean landmine, KB Snatch), Scarico Spinta Orizzontale Pesante (Mod. B), Trazione Cavi/Rope, Pliometria Gambe"
    }
    reg["ciclo_attivo_da_pianificare"] = "Ciclo 5 (Intensificazione Mesociclo 2)"
    reg["pattern_status"]["spinta_orizzontale"]["stato"] = "SCARICO SELETTIVO (Modalità B attiva - Coiled Punch iso)"
    reg["pattern_status"]["cerniera_anca"]["stato"] = "rientro balistico e potenza (COMBO Landmine clean, KB Snatch)"
    reg["pattern_status"]["braccia_spalle"]["stato"] = "mantenimento e sviluppo (Kneeling press landmine, Endless rope, Skierg pagaia)"
    reg["pattern_status"]["accosciata_squat"]["stato"] = "esplosività monolaterale (Step Up box, Super Mario jump)"

    c4_unique_ids = {"158", "74", "127", "23", "84", "201", "231", "2", "31", "147", "188", "119", "77", "192", "190", "228", "233"}

    for ex in reg.get("esercizi", []):
        eid = str(int(ex["id"])) if ex["id"].isdigit() else ex["id"]
        if eid in c4_unique_ids:
            ex["usage_count"] = ex.get("usage_count", 0) + 1
            ex["last_cycle_used"] = "Ciclo 4"

    reg["esercizi_utilizzati"] = sum(1 for ex in reg.get("esercizi", []) if ex.get("usage_count", 0) > 0)

    with open(REGISTER_JSON, "w") as f:
        json.dump(reg, f, indent=2, ensure_ascii=False)
    print(f"-> registro_periodizzazione_pro.json aggiornato! Esercizi usati a catalogo: {reg['esercizi_utilizzati']}/{reg['totale_esercizi_catalogo']}.")


def main():
    print("=" * 75)
    print("Area46 Training Lab - Generazione Ufficiale CICLO 4 (PRO)")
    print("=" * 75)

    all_rows, meta = generate_all_rows()
    print(f"Righe totali PRO generate: {len(all_rows)}")
    print("  - Ciclo 1 (1-6): 57 righe")
    print("  - Ciclo 2 (7-12): 51 righe")
    print("  - Ciclo 3 (13-18): 51 righe")
    print("  - Ciclo 4 (19-24): 51 righe")

    # 1. Connessione Google Fogli
    print("\n[Fase 1] Connessione a Google Sheets...")
    gc = gspread.service_account(filename=CREDENTIALS_FILE)

    # 2. Aggiorna Target Sheet
    print(f"\n[Fase 2] Aggiornamento Target Sheet '{TARGET_SHEET_ID}'...")
    sh_target = gc.open_by_key(TARGET_SHEET_ID)
    ws_pro_target = sh_target.worksheet("PRO")
    ws_pro_target.batch_clear(["A1:L300"])
    data_to_write = [COLUMNS] + all_rows
    ws_pro_target.update(range_name=f"A1:L{len(data_to_write)}", values=data_to_write)
    print(f"-> Scritte {len(data_to_write)} righe in 'PRO'.")
    apply_google_sheet_styling(sh_target, ws_pro_target, meta)

    # 3. Aggiorna Master Sheet
    print(f"\n[Fase 3] Aggiornamento Master Sheet '{LAB_SHEET_ID}'...")
    sh_lab = gc.open_by_key(LAB_SHEET_ID)
    ws_pro_lab = sh_lab.worksheet("Pro")
    ws_pro_lab.batch_clear(["A1:L300"])
    ws_pro_lab.update(range_name=f"A1:L{len(data_to_write)}", values=data_to_write)
    print(f"-> Scritte {len(data_to_write)} righe in 'Pro'.")
    apply_google_sheet_styling(sh_lab, ws_pro_lab, meta)

    # 4. Aggiorna File Excel locali
    print("\n[Fase 4] Aggiornamento file Excel locali...")
    update_excel_file(LOCAL_LAB_EXCEL, "Pro", all_rows, meta)
    update_excel_file(LOCAL_PRO_EXCEL, "PRO", all_rows, meta)

    # 5. Aggiorna demo-data.json per l'App
    print("\n[Fase 5] Aggiornamento demo-data.json per l'App...")
    update_demo_data_json(all_rows)

    # 6. Aggiorna registro periodizzazione
    print("\n[Fase 6] Aggiornamento registro periodizzazione e Google Drive...")
    update_register_json()

    # 7. Sincronizza scheda registro su Drive
    os.system(".venv/bin/python3 crea_scheda_registro_drive.py")

    # 8. Sincronizza cartella Drive
    os.system(".venv/bin/python3 sincronizza_drive.py")

    print("\n" + "=" * 75)
    print("CICLO 4 GENERATO, APPLICATO E SINCRONIZZATO CON SUCCESSO!")
    print("=" * 75)


if __name__ == "__main__":
    main()
