#!/usr/bin/env python3
"""
Generazione e Applicazione Ufficiale del CICLO 3 (Area46 Training Lab - Livello PRO).
Sincronizza:
- Google Sheet TARGET: 'PRO database esercizi' (scheda 'PRO')
- Google Sheet MASTER: 'database esercizi Lab' (scheda 'Pro')
- File Excel locali: 'PRO database esercizi.xlsx' e 'database esercizi Lab.xlsx'
- App Data: 'demo-data.json'
- Registro Periodizzazione: 'registro_periodizzazione_pro.json'
"""

import os
import json
import gspread
import openpyxl
from openpyxl.styles import PatternFill, Border, Side, Font, Alignment

CREDENTIALS_FILE = "credenziali.json"
TARGET_SHEET_ID = "1AQHkYjXhI6WhRPORJ4DntiGZQd8ob54wJxkwP_LymwY" # PRO database esercizi
LAB_SHEET_ID = "1SNy8IJk1E20BcAdar0F1TqqnpRIDkjlkNn6P2ljKPNY"    # database esercizi Lab
LOCAL_LAB_EXCEL = "database esercizi Lab.xlsx"
LOCAL_PRO_EXCEL = "PRO database esercizi.xlsx"
DEMO_DATA_JSON = "demo-data.json"
REGISTER_JSON = "registro_periodizzazione_pro.json"

COLUMNS = [
    "livello",
    "Settimana",
    "giorno",
    "Sequenza",
    "ID esercizio",
    "nome esercizio",
    "parametri",
    "recupero",
    "minutaggio blocco",
    "Note Tecniche",
    "video temp",
    "data_pubblicazione"
]

VIDEO_MAP = {
    "93": "https://youtu.be/aJ98KT13FtE?si=Li6fp-kbyh4wCnfT",
    "163": "https://youtu.be/WL2Gbf9klSI?si=R-4lfDqWZo7_yc9B",
    "79": "https://youtu.be/q_OfuLjIMFI",
    "78": "https://youtu.be/nSj7QKBbvWA?si=fGWg_soHjRYO0xex",
    "75": "https://youtu.be/V-5Q4l-SM20",
    "80": "https://youtu.be/2gCslk95r20",
    "234": "https://youtube.com/shorts/iSipY1JCKdQ",
    "34": "",
    "33": "https://youtube.com/shorts/GU3ePuvJwMk",
    "220": "https://youtu.be/EBv0BjKhUlA?si=I07YNW02JWJF-oiD",
    "218": "https://youtu.be/zIqsjVRtc5o",
    "82": "https://youtu.be/jqhow36DTdg",
    "191": "https://youtu.be/rRgKOg0qaA4?si=aYRTcGburOhwv1a6",
    "95": "https://youtube.com/shorts/arECYW5JFeA",
    "120": "https://youtube.com/shorts/0uI_9_0qR7k",
    "205": "https://youtu.be/OGEUeBvoKc0?si=v7LIiDqSnd71NyM4",
    "225": "https://youtu.be/Nwd-5fuLm3k?si=jGWxuH-RfvPTnkWN"
}

# --- CICLO 1 ---
C1_A = [
    ("1a", 51, "High hand switch", "3x30", "", "jump set"),
    ("1b", 54, "Landmine press Cross Step", "3x8+8", "", "jump set"),
    ("2",  56, "Landmine coiled Deadlift", "3x6+6", 'RBS 60"', "rpe 8"),
    ("3",  43, "Pounce squat", "4x8", '90"', "rpe 8"),
    ("4a", 82, "Landmine Biker's Squat", "3x10", "", "jump set"),
    ("4b", 81, "Meadow Row", "3x10+10", '60"', "jump set"),
    ("5a", 69, "Split Snatch", '5x30"', "", "circuito"),
    ("5b", 170, "Coiled Kettlebell Swing", '5x30"', "", "circuito"),
    ("5c", 100, "Speed skater touch", '5x30"', '45"', "circuito")
]

C1_B = [
    ("1a", 46, "Split Switch Screwdriver", "3x30", "", "jump set"),
    ("1b", 64, "High Pull ISO", '3x15"+15"', "", "jump set"),
    ("2a", 59, "Split Clean and Jerk", "3x6+6 (rpe 8)", 'RBS 45"', "jump set"),
    ("2b", 45, "Landmine rotational Lunges", "3x8 (RPE8)", '45"', "jump set"),
    ("3a", 230, "kettlebell clean and press", "3x6+6", 'RBS 45"', "jump set"),
    ("3b", 44, "Coiled Reverse Lunge", "3x6+6", 'RBS 45"', "jump set"),
    ("4a", 205, "Kick-through", '4x30"', '60"', "jump set"),
    ("4b", 206, "Side kick through", '4x30"', '60"', "jump set"),
    ("5a", 63, "Landmine Clean and jerk NO DOWN", "x 8+8", "", "AMRAP 10'"),
    ("5b", 203, "Plyometric lunges", "x 10", "", "AMRAP 10'")
]

# --- CICLO 2 ---
C2_A = [
    ("1a", 98, "Avenger Blade passing", "3x24", "", "jump set"),
    ("1b", 49, "Step/Step back press", "3x8+8", '45"', "jump set"),
    ("2",  79, "Bench press", "4x8", '90"', "RPE 8"),
    ("3a", 81, "Meadow Row", "4x8+8", '60"', "jump set | RPE 8"),
    ("3b", 82, "Landmine Biker's Squat", "4x8+8", 'RBS 60"', "jump set | RPE 8"),
    ("4a", 172, "Kettlebell swing to clean", "4x8+8", "", "circuito (4 round)"),
    ("4b", 124, "Slamball Skater jump", "4x10", "", "circuito (4 round)"),
    ("4c", 1,   "skierg regular", '4x40"', '60"', "circuito (4 round)")
]

C2_B = [
    ("1a", 55, "Hang Position ISO", '3x20"+20"', "", "jump set"),
    ("1b", 116, "High pull bounce", "3x10+10", '45"', "jump set"),
    ("2a", 53, "Split C&J Touch & Go", "4x6+6", 'RBS 45"', "jump set | RPE 8"),
    ("2b", 76, "Push the wall row", "4x8+8", 'RBS 45"', "jump set | RPE 8"),
    ("3a", 101, "Sprinter Push up", "4x10", '45"', "jump set | RPE 8"),
    ("3b", 44, "Coiled Reverse Lunge", "4x8+8", 'RBS 45"', "jump set | RPE 8"),
    ("4a", 73, "Lateral landmine clean", "x 6+6", "", "AMRAP 15'"),
    ("4b", 206, "Side kick through", "x 8+8", "", "AMRAP 15'"),
    ("4c", 187, "Sprawl to Broad Jump", "x 8", "", "AMRAP 15'")
]

# --- CICLO 3 (NUOVO CORSO METODOLOGICO) ---
# Obiettivi:
# - Scarico Selettivo Cerniera d'Anca (Modalità B: zero stacchi pesanti, sostituiti da mobilità/destrezza)
# - Focus Prioritario Braccia e Spalle (bicipiti, tricipiti, spalle, trazione/spinta)
# - Progressione Bench Press (da 4x8 a 4x6 @ RPE 8.5)
# - Inserimento Pull up Progressioni (ID 78), Landmine Jerk (ID 218), One Arm Press (ID 75), Endless Rope (ID 191)
# - Tetto 60' rigorosamente rispettato

C3_A = [
    ("1a", 93,  "Pallof Press con elastico", "3x10+10", "", "jump set"),
    ("1b", 163, "Arm basic DB pendulum", "3x10+10", '45"', "jump set"),
    ("2",  79,  "Bench press", "4x6", '90"', "RPE 8.5"),
    ("3a", 78,  "pull up Progressioni", "4x6-8", "", "jump set | RPE 8"),
    ("3b", 75,  "One arm press", "4x8+8", '60"', "jump set | RPE 8"),
    ("4a", 80,  "Coiled Bench Row", "3x10+10", "", "circuito (3 round)"),
    ("4b", 234, "rebound push up", "3x8", "", "circuito (3 round)"),
    ("4c", 34,  "air bike solo braccia", '3x30"', '60"', "circuito (3 round)")
]

C3_B = [
    ("1a", 33,  "Coiled Cobra", "3x8+8", "", "jump set"),
    ("1b", 220, "Wall High Pull", "3x10+10", '45"', "jump set"),
    ("2a", 218, "Landmine Jerk", "4x5+5", 'RBS 60"', "jump set | RPE 8.5"),
    ("2b", 82,  "Landmine Biker's Squat", "4x8", 'RBS 60"', "jump set | RPE 8"),
    ("3a", 191, "Endless rope Standing Power Pull", '4x20"', "", "jump set"),
    ("3b", 95,  "Yin yang sit up press", "4x8+8", '45"', "jump set"),
    ("4a", 120, "Aerial Jerk", "x 6+6", "", "AMRAP 12'"),
    ("4b", 205, "Kick-through progression", "x 10", "", "AMRAP 12'"),
    ("4c", 225, "Ball over shoulder", "x 8", "", "AMRAP 12'")
]


def generate_all_rows():
    """Genera le 159 righe totali per Ciclo 1, Ciclo 2 e Ciclo 3."""
    all_rows = []
    meta = []

    cicli_def = [
        ("Ciclo 1", [
            (1, "A", C1_A), (2, "B", C1_B),
            (3, "A", C1_A), (4, "B", C1_B),
            (5, "A", C1_A), (6, "B", C1_B)
        ]),
        ("Ciclo 2", [
            (7, "A", C2_A), (8, "B", C2_B),
            (9, "A", C2_A), (10, "B", C2_B),
            (11, "A", C2_A), (12, "B", C2_B)
        ]),
        ("Ciclo 3", [
            (13, "A", C3_A), (14, "B", C3_B),
            (15, "A", C3_A), (16, "B", C3_B),
            (17, "A", C3_A), (18, "B", C3_B)
        ])
    ]

    for ciclo_name, plan in cicli_def:
        last_sess_num = plan[-1][0]
        for sess_num, sess_type, exercises in plan:
            giorno_str = f"{sess_num} allenamento {sess_type}"
            is_a = (sess_type == "A")
            for i, ex in enumerate(exercises):
                is_last_in_cycle = (sess_num == last_sess_num and i == len(exercises) - 1)
                vid = VIDEO_MAP.get(str(ex[1]), "")
                row = [
                    "PRO",
                    ciclo_name,
                    giorno_str,
                    ex[0],
                    ex[1],
                    ex[2],
                    ex[3],
                    ex[4],
                    "",
                    ex[5],
                    vid,
                    ""
                ]
                all_rows.append(row)
                meta.append({"is_a": is_a, "is_cycle_end": is_last_in_cycle})

    return all_rows, meta


def apply_google_sheet_styling(spreadsheet, worksheet, meta):
    """Applica formattazione visiva con righe alternate A/B e bordi solidi di fine ciclo."""
    sheet_id = worksheet.id
    requests = []

    # 1. Header
    requests.append({
        "repeatCell": {
            "range": {
                "sheetId": sheet_id,
                "startRowIndex": 0,
                "endRowIndex": 1,
                "startColumnIndex": 0,
                "endColumnIndex": len(COLUMNS)
            },
            "cell": {
                "userEnteredFormat": {
                    "textFormat": {"bold": True},
                    "backgroundColor": {"red": 0.88, "green": 0.88, "blue": 0.88}
                }
            },
            "fields": "userEnteredFormat(textFormat,backgroundColor)"
        }
    })

    # 2. Alternanza colori
    idx = 0
    while idx < len(meta):
        is_a = meta[idx]["is_a"]
        start_row = idx + 1
        while idx < len(meta) and meta[idx]["is_a"] == is_a:
            idx += 1
        end_row = idx + 1

        bg_color = {"red": 0.95, "green": 0.95, "blue": 0.95} if is_a else {"red": 1.0, "green": 1.0, "blue": 1.0}
        requests.append({
            "repeatCell": {
                "range": {
                    "sheetId": sheet_id,
                    "startRowIndex": start_row,
                    "endRowIndex": end_row,
                    "startColumnIndex": 0,
                    "endColumnIndex": len(COLUMNS)
                },
                "cell": {
                    "userEnteredFormat": {
                        "backgroundColor": bg_color
                    }
                },
                "fields": "userEnteredFormat.backgroundColor"
            }
        })

    # 3. Bordi di fine ciclo
    for idx, m in enumerate(meta):
        if m["is_cycle_end"]:
            row_idx = idx + 1
            requests.append({
                "updateBorders": {
                    "range": {
                        "sheetId": sheet_id,
                        "startRowIndex": row_idx,
                        "endRowIndex": row_idx + 1,
                        "startColumnIndex": 0,
                        "endColumnIndex": len(COLUMNS)
                    },
                    "bottom": {
                        "style": "SOLID_MEDIUM",
                        "width": 2,
                        "color": {"red": 0.0, "green": 0.0, "blue": 0.0}
                    }
                }
            })

    body = {"requests": requests}
    spreadsheet.batch_update(body)
    print(f"-> Styling applicato su '{worksheet.title}'.")


def update_excel_file(filepath, sheet_name, all_rows, meta):
    """Aggiorna e formatta con openpyxl."""
    if not os.path.exists(filepath):
        return

    wb = openpyxl.load_workbook(filepath)
    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(title=sheet_name)
    else:
        ws = wb[sheet_name]

    # Svuota
    for r in range(1, max(ws.max_row + 1, 250)):
        for c in range(1, 15):
            cell = ws.cell(row=r, column=c)
            cell.value = None
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border()

    # Header
    header_fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
    header_font = Font(bold=True)
    for col_idx, col_name in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font

    fill_a = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    fill_b = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    border_bottom_thick = Border(bottom=Side(style="medium", color="000000"))

    for row_idx, (r_vals, m) in enumerate(zip(all_rows, meta), 2):
        row_fill = fill_a if m["is_a"] else fill_b
        is_end = m["is_cycle_end"]
        for col_idx, val in enumerate(r_vals, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = row_fill
            if is_end:
                cell.border = border_bottom_thick

    wb.save(filepath)
    print(f"-> File Excel locale '{filepath}' (foglio '{sheet_name}') aggiornato!")


def update_demo_data_json(all_rows):
    """Sincronizza demo-data.json con tutte le righe PRO (Ciclo 1, 2 e 3)."""
    if not os.path.exists(DEMO_DATA_JSON):
        return

    with open(DEMO_DATA_JSON, "r") as f:
        demo_data = json.load(f)

    # Rimuovi vecchie righe PRO
    demo_data["allenamenti"] = [a for a in demo_data["allenamenti"] if a.get("livello") != "PRO"]

    existing_ids = [a.get("id", 0) for a in demo_data["allenamenti"]]
    max_id = max(existing_ids) if existing_ids else 500

    for r in all_rows:
        max_id += 1
        demo_data["allenamenti"].append({
            "id": max_id,
            "livello": r[0],
            "settimana": r[1],
            "giorno": r[2],
            "sequenza": r[3],
            "id_esercizio": f"{int(r[4]):03d}",
            "nome_esercizio": r[5],
            "parametri": r[6],
            "recupero": r[7],
            "minutaggio_blocco": "",
            "note_tecniche": r[9],
            "link_video": r[10],
            "data_pubblicazione": None
        })

    with open(DEMO_DATA_JSON, "w") as f:
        json.dump(demo_data, f, indent=2, ensure_ascii=False)
    print(f"-> demo-data.json aggiornato con {len(all_rows)} righe PRO (Cicli 1, 2, 3).")


def update_register_json():
    """Aggiorna registro_periodizzazione_pro.json con il Ciclo 3."""
    if not os.path.exists(REGISTER_JSON):
        return

    with open(REGISTER_JSON, "r") as f:
        reg = json.load(f)

    reg["storico_cicli"]["Ciclo 3"] = {
        "status": "attivo / programmato",
        "sessioni": 6,
        "focus": "Scarico anca (Mod. B), Priorità Braccia/Spalle, Bench press progressione 4x6, Pull up, Landmine Jerk"
    }
    reg["ciclo_attivo_da_pianificare"] = "Ciclo 4 (Completamento Mesociclo 1 / Deload selettivo)"
    reg["pattern_status"]["spinta_orizzontale"]["stato"] = "picco di forza (Bench press 4x6 @ RPE 8.5)"
    reg["pattern_status"]["cerniera_anca"]["stato"] = "SCARICO SELETTIVO (Modalità B attiva)"
    reg["pattern_status"]["braccia_spalle"]["stato"] = "attivato e prioritario (One arm press, Landmine jerk, Endless rope, Air bike braccia)"

    c3_unique_ids = {"93", "163", "79", "78", "75", "80", "234", "34", "33", "220", "218", "82", "191", "95", "120", "205", "225"}

    for ex in reg.get("esercizi", []):
        eid = str(int(ex["id"])) if ex["id"].isdigit() else ex["id"]
        if eid in c3_unique_ids:
            ex["usage_count"] = ex.get("usage_count", 0) + 1
            ex["last_cycle_used"] = "Ciclo 3"

    reg["esercizi_utilizzati"] = sum(1 for ex in reg.get("esercizi", []) if ex.get("usage_count", 0) > 0)

    with open(REGISTER_JSON, "w") as f:
        json.dump(reg, f, indent=2, ensure_ascii=False)
    print(f"-> registro_periodizzazione_pro.json aggiornato! Esercizi usati a catalogo: {reg['esercizi_utilizzati']}/{reg['totale_esercizi_catalogo']}.")


def main():
    print("=" * 75)
    print("Area46 Training Lab - Generazione Ufficiale CICLO 3 (PRO)")
    print("=" * 75)

    all_rows, meta = generate_all_rows()
    print(f"Righe totali PRO generate: {len(all_rows)}")
    print("  - Ciclo 1 (1-6): 57 righe")
    print("  - Ciclo 2 (7-12): 51 righe")
    print("  - Ciclo 3 (13-18): 51 righe")

    # 1. Connessione Google Fogli
    print("\n[Fase 1] Connessione a Google Sheets...")
    gc = gspread.service_account(filename=CREDENTIALS_FILE)

    # 2. Aggiorna Target Sheet
    print(f"\n[Fase 2] Aggiornamento Target Sheet '{TARGET_SHEET_ID}'...")
    sh_target = gc.open_by_key(TARGET_SHEET_ID)
    ws_pro_target = sh_target.worksheet("PRO")
    ws_pro_target.batch_clear(["A1:L250"])
    data_to_write = [COLUMNS] + all_rows
    ws_pro_target.update(range_name=f"A1:L{len(data_to_write)}", values=data_to_write)
    print(f"-> Scritte {len(data_to_write)} righe in 'PRO'.")
    apply_google_sheet_styling(sh_target, ws_pro_target, meta)

    # 3. Aggiorna Master Sheet
    print(f"\n[Fase 3] Aggiornamento Master Sheet '{LAB_SHEET_ID}'...")
    sh_lab = gc.open_by_key(LAB_SHEET_ID)
    ws_pro_lab = sh_lab.worksheet("Pro")
    ws_pro_lab.batch_clear(["A1:L250"])
    ws_pro_lab.update(range_name=f"A1:L{len(data_to_write)}", values=data_to_write)
    print(f"-> Scritte {len(data_to_write)} righe in 'Pro'.")
    apply_google_sheet_styling(sh_lab, ws_pro_lab, meta)

    # 4. Aggiorna File Excel locali
    print("\n[Fase 4] Aggiornamento file Excel locali...")
    update_excel_file(LOCAL_LAB_EXCEL, "Pro", all_rows, meta)
    update_excel_file(LOCAL_PRO_EXCEL, "PRO", all_rows, meta)

    # 5. Aggiorna demo-data.json per l'App
    print("\n[Fase 5] Aggiornamento demo-data.json per l'App...")
    update_demo_data_json(all_rows)

    # 6. Aggiorna registro periodizzazione
    print("\n[Fase 6] Aggiornamento registro periodizzazione e Google Drive...")
    update_register_json()

    # 7. Sincronizza scheda registro su Drive
    os.system(".venv/bin/python3 crea_scheda_registro_drive.py")

    # 8. Sincronizza cartella Drive
    os.system(".venv/bin/python3 sincronizza_drive.py")

    print("\n" + "=" * 75)
    print("CICLO 3 GENERATO, APPLICATO E SINCRONIZZATO CON SUCCESSO!")
    print("=" * 75)


if __name__ == "__main__":
    main()
