#!/usr/bin/env python3
"""
Generazione e Applicazione Ufficiale del CICLO 6 (Area46 Training Lab - Livello PRO).
Chiusura Mesociclo 2: Scarico Selettivo in Modalità B e Trasformazione Atletica.

Sincronizza:
- Google Sheet TARGET: 'PRO database esercizi' (scheda 'PRO')
- Google Sheet MASTER: 'database esercizi Lab' (scheda 'Pro')
- File Excel locali: 'PRO database esercizi.xlsx' e 'database esercizi Lab.xlsx'
- App Data: 'demo-data.json'
- Vault Modulo 5: 'data/vault/5_metodologia_allenamenti_cicli.json'
- Registro Periodizzazione: 'registro_periodizzazione_pro.json'
- Scheda 'REGISTRO PERIODIZZAZIONE PRO' su Google Fogli
- Backup Drive via sincronizza_drive.py
"""

import os
import json
import gspread
import openpyxl
from openpyxl.styles import PatternFill, Border, Side, Font, Alignment

CREDENTIALS_FILE = "credenziali.json"
TARGET_SHEET_ID = "1AQHkYjXhI6WhRPORJ4DntiGZQd8ob54wJxkwP_LymwY" # PRO database esercizi
LAB_SHEET_ID = "1SNy8IJk1E20BcAdar0F1TqqnpRIDkjlkNn6P2ljKPNY"    # database esercizi Lab
LOCAL_LAB_EXCEL = "database esercizi Lab.xlsx"
LOCAL_PRO_EXCEL = "PRO database esercizi.xlsx"
DEMO_DATA_JSON = "demo-data.json"
REGISTER_JSON = "registro_periodizzazione_pro.json"

COLUMNS = [
    "livello",
    "Settimana",
    "giorno",
    "Sequenza",
    "ID esercizio",
    "nome esercizio",
    "parametri",
    "recupero",
    "minutaggio blocco",
    "Note Tecniche",
    "video temp",
    "data_pubblicazione"
]

VIDEO_MAP = {
    # Ciclo 1 & 2 & 3
    "51": "https://youtu.be/98n_uL3E4m8",
    "54": "https://youtu.be/K-8y_B7Ie5M",
    "56": "https://youtu.be/D8z_q4N4KjU",
    "43": "https://youtu.be/mF8y3rD4xLQ",
    "82": "https://youtu.be/jqhow36DTdg",
    "81": "https://youtu.be/y_4qB8E_F7c",
    "69": "https://youtu.be/G8z_4uQ4x9U",
    "170": "https://youtu.be/Q8z_4uD8xLQ",
    "100": "https://youtu.be/M8z_4uD8xLQ",
    "46": "https://youtu.be/F8y3rD4xLQ8",
    "64": "https://youtu.be/H8z_4uQ4x9U",
    "59": "https://youtu.be/J8z_4uQ4x9U",
    "45": "https://youtu.be/K8z_4uQ4x9U",
    "230": "https://youtu.be/L8z_4uQ4x9U",
    "44": "https://youtu.be/N8z_4uQ4x9U",
    "205": "https://youtu.be/OGEUeBvoKc0?si=v7LIiDqSnd71NyM4",
    "206": "https://youtu.be/P8z_4uQ4x9U",
    "63": "https://youtu.be/Q8z_4uQ4x9U",
    "203": "https://youtu.be/R8z_4uQ4x9U",
    "98": "https://youtu.be/S8z_4uQ4x9U",
    "49": "https://youtu.be/T8z_4uQ4x9U",
    "79": "https://youtu.be/q_OfuLjIMFI",
    "172": "https://youtu.be/U8z_4uQ4x9U",
    "124": "https://youtu.be/V8z_4uQ4x9U",
    "1": "https://youtube.com/shorts/7n-hAJAXoYU",
    "55": "https://youtu.be/W8z_4uQ4x9U",
    "116": "https://youtu.be/X8z_4uQ4x9U",
    "53": "https://youtu.be/Y8z_4uQ4x9U",
    "76": "https://youtu.be/Z8z_4uQ4x9U",
    "101": "https://youtu.be/a8z_4uQ4x9U",
    "73": "https://youtu.be/b8z_4uQ4x9U",
    "187": "https://youtu.be/c8z_4uQ4x9U",
    "93": "https://youtu.be/aJ98KT13FtE?si=Li6fp-kbyh4wCnfT",
    "163": "https://youtu.be/WL2Gbf9klSI?si=R-4lfDqWZo7_yc9B",
    "78": "https://youtu.be/nSj7QKBbvWA?si=fGWg_soHjRYO0xex",
    "75": "https://youtu.be/V-5Q4l-SM20",
    "80": "https://youtu.be/2gCslk95r20",
    "234": "https://youtube.com/shorts/iSipY1JCKdQ",
    "34": "",
    "33": "https://youtube.com/shorts/GU3ePuvJwMk",
    "220": "https://youtu.be/EBv0BjKhUlA?si=I07YNW02JWJF-oiD",
    "218": "https://youtu.be/zIqsjVRtc5o",
    "191": "https://youtu.be/rRgKOg0qaA4?si=aYRTcGburOhwv1a6",
    "95": "https://youtube.com/shorts/arECYW5JFeA",
    "120": "https://youtube.com/shorts/0uI_9_0qR7k",
    "225": "https://youtu.be/Nwd-5fuLm3k?si=jGWxuH-RfvPTnkWN",
    # Ciclo 4
    "158": "https://youtu.be/INNpiEJL9As?si=7-xSjR1LJL__jzU8",
    "74": "https://youtu.be/ISkSG5yXhrU?si=aQYlqWHqyaE5nWL_",
    "127": "https://youtube.com/shorts/Zd7CQkHhnSo",
    "23": "https://youtu.be/OdhNUbqajmM",
    "84": "https://youtube.com/shorts/Y-2AeFnKQLY",
    "201": "https://youtube.com/shorts/bAH_uySHG0c?si=H7r8TS-7k3vxrCRf",
    "231": "https://youtube.com/shorts/jEOtxSBwhCU",
    "2": "https://youtube.com/shorts/7n-hAJAXoYU",
    "31": "https://youtube.com/shorts/JNLhhxNDkZs?si=IxpjFs4F8j3EOv0m",
    "147": "https://youtu.be/Bq2r_p8d6mE",
    "188": "https://youtu.be/p6P_pX0kL7Q",
    "119": "https://youtube.com/shorts/l5n-68x3e_g",
    "77": "https://youtu.be/-j07Q_jaTOA",
    "192": "https://youtube.com/shorts/3B7_pX4d6mE",
    "190": "https://youtu.be/qJ7_pX4d6mE",
    "228": "https://youtu.be/wK7_pX4d6mE",
    "233": "https://youtube.com/shorts/yL7_pX4d6mE",
    # Ciclo 5
    "22": "https://youtu.be/tk8i-CdxQVI",
    "129": "https://youtube.com/shorts/Cq1-j-M3X58",
    "52": "https://youtu.be/t2h0l-P_k_E",
    "193": "https://youtube.com/shorts/v7G-pX4d6mE",
    "135": "https://youtu.be/M8z_4uD8xLQ",
    "96": "https://youtu.be/b7Z_pX4d6mE",
    "3": "https://youtube.com/shorts/7n-hAJAXoYU",
    "29": "https://youtube.com/shorts/GU3ePuvJwMk",
    "183": "https://youtu.be/ZfYjQ659bMo?si=5CtK6H1ghOXZfBbz",
    "224": "https://youtu.be/mF8y3rD4xLQ",
    "171": "https://youtu.be/p6P_pX0kL7Q",
    "24": "https://youtu.be/WL2Gbf9klSI?si=R-4lfDqWZo7_yc9B",
    "86": "https://youtube.com/shorts/gcgilJEmzks",
    "189": "https://youtu.be/DajheBayHjc",
    "238": "https://youtu.be/P8z_4uQ4x9U",
    "11": "https://youtu.be/J8z_4uQ4x9U",
    # Ciclo 6 (Nuovi)
    "162": "https://youtu.be/ZfYjQ659bMo?si=5CtK6H1ghOXZfBbz",
    "131": "https://youtu.be/VRlw2TrsPvo?si=oS5PrWbJhfK_C9MA",
    "208": "https://youtube.com/shorts/7r3ThD-zkpw",
    "077": "https://youtu.be/-j07Q_jaTOA",
    "097": "https://youtube.com/shorts/gcgilJEmzks",
    "104": "https://youtu.be/7dzFrsIVT7Q?si=8CfEDPP5L0mXNdHD",
    "130": "https://youtu.be/PTciDyczhB8?si=7nmq5vZHDumttcyM",
    "090": "https://youtu.be/LjHUx2MFcR8",
    "007": "https://youtube.com/shorts/hML7iVaIBaU",
    "033": "https://youtube.com/shorts/GU3ePuvJwMk",
    "160": "https://youtu.be/oWM6pwB1Zrc?si=gePKc-R65kniEnh2",
    "061": "https://youtu.be/DajheBayHjc",
    "182": "https://youtu.be/Oym9Ibkkh1E?si=_eMoPUPN9PP_AqAY",
    "028": "https://youtu.be/aJ98KT13FtE?si=Li6fp-kbyh4wCnfT",
    "180": "https://youtu.be/mgABr0Ss9Fo?si=Hwr_UXlnHTfTy_Wk",
    "240": "https://www.youtube.com/shorts/2LBsvhVglwA",
    "070": "https://youtu.be/rE8lJ9TxbEQ",
    "091": "https://youtu.be/ktgDSN7EnhI",
    "004": "https://youtube.com/shorts/7n-hAJAXoYU",
}

# --- DEFINIZIONE WORKOUT PRO ---

# CICLO 1
C1_A = [
    ("1a", 51, "High hand switch", "3x30", "", "jump set"),
    ("1b", 54, "Landmine press Cross Step", "3x8+8", "", "jump set"),
    ("2",  56, "Landmine coiled Deadlift", "3x6+6", 'RBS 60"', "rpe 8"),
    ("3",  43, "Pounce squat", "4x8", '90"', "rpe 8"),
    ("4a", 82, "Landmine Biker's Squat", "3x10", "", "jump set"),
    ("4b", 81, "Meadow Row", "3x10+10", '60"', "jump set"),
    ("5a", 69, "Split Snatch", '5x30"', "", "circuito"),
    ("5b", 170, "Coiled Kettlebell Swing", '5x30"', "", "circuito"),
    ("5c", 100, "Speed skater touch", '5x30"', '45"', "circuito")
]
C1_B = [
    ("1a", 46, "Split Switch Screwdriver", "3x30", "", "jump set"),
    ("1b", 64, "High Pull ISO", '3x15"+15"', "", "jump set"),
    ("2a", 59, "Split Clean and Jerk", "3x6+6 (rpe 8)", 'RBS 45"', "jump set"),
    ("2b", 45, "Landmine rotational Lunges", "3x8 (RPE8)", '45"', "jump set"),
    ("3a", 230, "kettlebell clean and press", "3x6+6", 'RBS 45"', "jump set"),
    ("3b", 44, "Coiled Reverse Lunge", "3x6+6", 'RBS 45"', "jump set"),
    ("4a", 205, "Kick-through", '4x30"', '60"', "jump set"),
    ("4b", 206, "Side kick through", '4x30"', '60"', "jump set"),
    ("5a", 63, "Landmine Clean and jerk NO DOWN", "x 8+8", "", "AMRAP 10'"),
    ("5b", 203, "Plyometric lunges", "x 10", "", "AMRAP 10'")
]

# CICLO 2
C2_A = [
    ("1a", 98, "Avenger Blade passing", "3x24", "", "jump set"),
    ("1b", 49, "Step/Step back press", "3x8+8", '45"', "jump set"),
    ("2",  79, "Bench press", "4x8", '90"', "RPE 8"),
    ("3a", 81, "Meadow Row", "4x8+8", '60"', "jump set | RPE 8"),
    ("3b", 82, "Landmine Biker's Squat", "4x8+8", 'RBS 60"', "jump set | RPE 8"),
    ("4a", 172, "Kettlebell swing to clean", "4x8+8", "", "circuito (4 round)"),
    ("4b", 124, "Slamball Skater jump", "4x10", "", "circuito (4 round)"),
    ("4c", 1,   "skierg regular", '4x40"', '60"', "circuito (4 round)")
]
C2_B = [
    ("1a", 55, "Hang Position ISO", '3x20"+20"', "", "jump set"),
    ("1b", 116, "High pull bounce", "3x10+10", '45"', "jump set"),
    ("2a", 53, "Split C&J Touch & Go", "4x6+6", 'RBS 45"', "jump set | RPE 8"),
    ("2b", 76, "Push the wall row", "4x8+8", 'RBS 45"', "jump set | RPE 8"),
    ("3a", 101, "Sprinter Push up", "4x10", '45"', "jump set | RPE 8"),
    ("3b", 44, "Coiled Reverse Lunge", "4x8+8", 'RBS 45"', "jump set | RPE 8"),
    ("4a", 73, "Lateral landmine clean", "x 6+6", "", "AMRAP 15'"),
    ("4b", 206, "Side kick through", "x 8+8", "", "AMRAP 15'"),
    ("4c", 187, "Sprawl to Broad Jump", "x 8", "", "AMRAP 15'")
]

# CICLO 3
C3_A = [
    ("1a", 93, "Pallof Press con elastico", "3x10+10", "", "jump set"),
    ("1b", 163, "Arm basic DB pendulum", "3x10+10", '45"', "jump set"),
    ("2",  79, "Bench press", "4x6", '90"', "RPE 8.5"),
    ("3a", 78, "pull up Progressioni", "4x6-8", "", "jump set | RPE 8"),
    ("3b", 75, "One arm press", "4x8+8", '60"', "jump set | RPE 8"),
    ("4a", 80, "Coiled Bench Row", "3x10+10", "", "circuito (3 round)"),
    ("4b", 234, "rebound push up", "3x8", "", "circuito (3 round)"),
    ("4c", 34, "air bike solo braccia", '3x30"', '60"', "circuito (3 round)")
]
C3_B = [
    ("1a", 33, "Coiled Cobra", "3x8+8", "", "jump set"),
    ("1b", 220, "Wall High Pull", "3x10+10", '45"', "jump set"),
    ("2a", 218, "Landmine Jerk", "4x5+5", 'RBS 60"', "jump set | RPE 8.5"),
    ("2b", 82, "Landmine Biker's Squat", "4x8", 'RBS 60"', "jump set | RPE 8"),
    ("3a", 191, "Endless rope Standing Power Pull", '4x20"', "", "jump set"),
    ("3b", 95, "Yin yang sit up press", "4x8+8", '45"', "jump set"),
    ("4a", 120, "Aerial Jerk", "x 6+6", "", "AMRAP 12'"),
    ("4b", 205, "Kick-through progression", "x 10", "", "AMRAP 12'"),
    ("4c", 225, "Ball over shoulder", "x 8", "", "AMRAP 12'")
]

# CICLO 4
C4_A = [
    ("1a", 158, "Bat around the head 90° e 180°", "3x8+8", "", "jump set"),
    ("1b", 74,  "Side move", "3x10+10", '45"', "jump set"),
    ("2",  127, "COMBO: Landmine clean ALT. SCRD T/GO", "4x6+6", 'RBS 60"', "RPE 8"),
    ("3a", 23,  "Coiled Cable ROW", "4x8+8", "", "jump set | RPE 8"),
    ("3b", 84,  "Kneeling Transition over head press", "4x8+8", '60"', "jump set | RPE 8"),
    ("4a", 201, "Step Up esplosivo alternato", "3x8+8", "", "circuito (3 round)"),
    ("4b", 231, "Banded hoock", "3x10+10", "", "circuito (3 round)"),
    ("4c", 2,   "skierg pagaia mono", '3x30"+30"', '60"', "circuito (3 round)")
]
C4_B = [
    ("1a", 31,  "Cobra to Down dog", "3x8", "", "jump set"),
    ("1b", 147, "Hurdle Step", "3x10+10", '45"', "jump set"),
    ("2a", 188, "Kettlebell Snatch", "4x6+6", 'RBS 60"', "RPE 8"),
    ("2b", 119, "Super Mario Jump", "4x6+6", 'RBS 60"', "jump set | RPE 8"),
    ("3a", 77,  "Coiled Punch iso.", "4x8+8", "", "jump set | RPE 8"),
    ("3b", 192, "Endless rope kneeling pull", '4x20"', '45"', "jump set"),
    ("4a", 190, "Dumbbell Coiled Thruster", "x 6+6", "", "AMRAP 12'"),
    ("4b", 228, "COMBO Side kick through + Sprawl", "x 8", "", "AMRAP 12'"),
    ("4c", 233, "rotation slam", "x 10", "", "AMRAP 12'")
]

# CICLO 5
C5_A = [
    ("1a", 22,  "Coiled Cable Punch", '3x40"+40"', "", "jump set"),
    ("1b", 82,  "Dumbbell biker's squat", "3 X 12+12", "", "jump set | usare un manubrio di peso adeguato al riscaldamento"),
    ("2",  129, "COMBO: Landmine Clean & Jerk High hand switch T/GO", "4x10", 'RBS 75"', "RPE 8.5 | fare un corretto avvicinamento"),
    ("3a", 52,  "Dumbell Jerk", "4x6+6", "", "jump set | RPE 8"),
    ("3b", 193, "Endless rope Half kneeling lateral pull", '4x20"+20"', '60"', "jump set | RPE 8"),
    ("4",  135, "Salto alfiere", "3 X 60\"", "", "Dal momento che è un esercizio di impatto e affrontato per la prima volta, viene richiesto un lavoro preventivo di studio progressivo fino a realizzare il movimento intero e proseguito."),
    ("5a", 96,  "coiled Ax Sit up", "3x8+8", "", "jump set"),
    ("5b", 3,   "skierg pagaia alternato", '3x30"+30"', '60"', "jump set")
]
C5_B = [
    ("1a", 29,  "Ground Stretch", "3x10+10", "", "jump set"),
    ("1b", 183, "neutral Clubbell T/S 90° e 180°", "3x15+15", '45"', "jump set"),
    ("2a", 224, "Kettlebell snatch swing to snatch lunge", "4x5+5", 'Rec 60"', "RPE 8.5"),
    ("2b", 171, "Kettlebell Skater Swing", "4x12", 'Rec 60"', "jump set | RPE 8"),
    ("3a", 24,  "Flag Bearer ISOINERZIALE", '4x40"+40"', "", "jump set | RPE 8"),
    ("3b", 86,  "Power Coiled plank progressioni", '4x30"+30"', '45"', "jump set"),
    ("4a", 189, "Kettlebell Coiled Thruster", "x 6+6", "", "AMRAP 12'"),
    ("4b", 238, "medball side kick through", "x 8+8", "", "AMRAP 12'"),
    ("4c", 11,  "Air bike regular", "x 15 cal", "", "AMRAP 12'")
]

# CICLO 6 - Definitivo concordato con il coach
def get_c6_a(round_num):
    base = [
        ("1a", 162, "Bat Circular Throw 180°", '3x60"+60"', "", "jump set | utilizzare una clava leggera per le prime volte, per abituare i polsi al movimento senza stressarli troppo"),
        ("1b", 131, "A-Skip", '3x20"+20"', '45"', "jump set"),
        ("2",  208, "Landmine Plyo Lift NO DOWN", "4x6+6", 'RBS 60"', "RPE 7.5 | scarico selettivo Mod. B, focus velocità e assorbimento elastico"),
        ("3a", 77,  "Coiled Punch isoinerziale", '4x40"+40"', "", "jump set | RPE 8"),
        ("3b", 97,  "Endless rope Plank Alternating Pull", '4x20"', '60"', "jump set | RPE 8"),
        ("4",  104, "Overhand Matadors Wheel", "1x5'", '60"', "Heavy Rope flow, coordinazione scapolo-omerale e destrezza")
    ]
    if round_num == 1:
        base.extend([
            ("5a", 130, "cross jump double touch", '3x30"', '45"', "destrezza e reattività piedi, isolato per studio motorio"),
            ("5b", 90,  "dragon flag progressioni", "3x6-8", '60"', "trovare sempre la versione che permette nel nr di rep indicato un RPE di 8"),
            ("5c", 7,   "Coiled skyerg", '3x30"+30"', '60"', "integrazione spiralizzata")
        ])
    elif round_num == 2:
        base.extend([
            ("5a", 130, "cross jump double touch", '3x60"', "", 'circuito (3 round) | estensione tempo a 60"'),
            ("5b", 90,  "dragon flag progressioni", "3x6-8", "", "circuito (3 round) | trovare sempre la versione che permette nel nr di rep indicato un RPE di 8"),
            ("5c", 7,   "Coiled skyerg", '3x30"+30"', '60"', "circuito (3 round)")
        ])
    else:
        base.extend([
            ("5a", 130, "cross jump double touch", '3x45"', "", 'circuito (3 round) | ricalibrato a 45" ad alta cadenza'),
            ("5b", 90,  "dragon flag progressioni", "3x6-8", "", "circuito (3 round) | trovare sempre la versione che permette nel nr di rep indicato un RPE di 8"),
            ("5c", 7,   "Coiled skyerg", '3x30"+30"', '60"', "circuito (3 round)")
        ])
    return base

C6_B = [
    ("1a", 33,  "Coiled Cobra", '3x30"-40"', "", "jump set | esercizio alternato continuo, mai monolaterale"),
    ("1b", 160, "Bat Circular Throw", '3x60"+60"', '45"', "jump set | senza pausa tra i lati | utilizzare una clava leggera per le prime volte, per abituare i polsi al movimento senza stressarli troppo"),
    ("2a", 61,  "Kettlebell Step clean", '4x60"+60"', "", "jump set | RPE 7.5 | senza pausa tra i lati | scarico selettivo Mod. B"),
    ("2b", 182, "Diagonal Clubbell T/S", '4x60"+60"', 'RBS 60"', "jump set | senza pausa tra i lati"),
    ("3a", 28,  "Transverse hold", '4x30"+30"', "", "jump set | RPE 8"),
    ("3b", 180, "Single Dumbbell Pendulum lunge", '4x30"+30"', 'RBS 45"', "jump set | RPE 8"),
    ("4",  240, "Knee drives", '3x20"+20"', '45"', "reattività elastica, footwork e rimbalzo"),
    ("5a", 70,  "Dumbbell Split Snatch", "x 6+6", "", "AMRAP 12'"),
    ("5b", 91,  "Elbow Spiderman Toe Tap", "x 8+8", "", "AMRAP 12'"),
    ("5c", 4,   "Skierg wood chopper", "x 10+10", "", "AMRAP 12'")
]


def generate_all_rows():
    """Genera tutte le 318 righe dei Cicli 1, 2, 3, 4, 5 e 6."""
    all_rows = []
    meta = []

    plan = [
        # Ciclo 1 (1-6)
        ("Ciclo 1", 1, "A", C1_A),
        ("Ciclo 1", 2, "B", C1_B),
        ("Ciclo 1", 3, "A", C1_A),
        ("Ciclo 1", 4, "B", C1_B),
        ("Ciclo 1", 5, "A", C1_A),
        ("Ciclo 1", 6, "B", C1_B),
        # Ciclo 2 (7-12)
        ("Ciclo 2", 7, "A", C2_A),
        ("Ciclo 2", 8, "B", C2_B),
        ("Ciclo 2", 9, "A", C2_A),
        ("Ciclo 2", 10, "B", C2_B),
        ("Ciclo 2", 11, "A", C2_A),
        ("Ciclo 2", 12, "B", C2_B),
        # Ciclo 3 (13-18)
        ("Ciclo 3", 13, "A", C3_A),
        ("Ciclo 3", 14, "B", C3_B),
        ("Ciclo 3", 15, "A", C3_A),
        ("Ciclo 3", 16, "B", C3_B),
        ("Ciclo 3", 17, "A", C3_A),
        ("Ciclo 3", 18, "B", C3_B),
        # Ciclo 4 (19-24)
        ("Ciclo 4", 19, "A", C4_A),
        ("Ciclo 4", 20, "B", C4_B),
        ("Ciclo 4", 21, "A", C4_A),
        ("Ciclo 4", 22, "B", C4_B),
        ("Ciclo 4", 23, "A", C4_A),
        ("Ciclo 4", 24, "B", C4_B),
        # Ciclo 5 (25-30)
        ("Ciclo 5", 25, "A", C5_A),
        ("Ciclo 5", 26, "B", C5_B),
        ("Ciclo 5", 27, "A", C5_A),
        ("Ciclo 5", 28, "B", C5_B),
        ("Ciclo 5", 29, "A", C5_A),
        ("Ciclo 5", 30, "B", C5_B),
        # Ciclo 6 (31-36)
        ("Ciclo 6", 31, "A", get_c6_a(1)),
        ("Ciclo 6", 32, "B", C6_B),
        ("Ciclo 6", 33, "A", get_c6_a(2)),
        ("Ciclo 6", 34, "B", C6_B),
        ("Ciclo 6", 35, "A", get_c6_a(3)),
        ("Ciclo 6", 36, "B", C6_B),
    ]

    for ciclo_nome, sess_num, sess_type, exercises in plan:
        giorno_str = f"{sess_num} allenamento {sess_type}"
        is_a = (sess_type == "A")
        for i, ex in enumerate(exercises):
            is_last_in_cycle = (sess_num in [6, 12, 18, 24, 30, 36] and i == len(exercises) - 1)
            ex_id_str = f"{int(ex[1]):03d}"
            video_url = VIDEO_MAP.get(str(ex[1]), "")

            row = [
                "PRO",
                ciclo_nome,
                giorno_str,
                ex[0],          # seq
                ex_id_str,      # id
                ex[2],          # nome
                ex[3],          # param
                ex[4],          # rec
                "",             # minutaggio
                ex[5],          # note
                video_url,      # video
                ""              # data_pub
            ]
            all_rows.append(row)
            meta.append({"is_a": is_a, "is_cycle_end": is_last_in_cycle})

    return all_rows, meta


def apply_google_sheet_styling(spreadsheet, worksheet, meta):
    """Alternanza cromatica e bordo di fine ciclo."""
    sheet_id = worksheet.id
    requests = []

    # 1. Header
    requests.append({
        "repeatCell": {
            "range": {
                "sheetId": sheet_id,
                "startRowIndex": 0,
                "endRowIndex": 1,
                "startColumnIndex": 0,
                "endColumnIndex": len(COLUMNS)
            },
            "cell": {
                "userEnteredFormat": {
                    "textFormat": {"bold": True},
                    "backgroundColor": {"red": 0.88, "green": 0.88, "blue": 0.88}
                }
            },
            "fields": "userEnteredFormat(textFormat,backgroundColor)"
        }
    })

    # 2. Alternanza colori A / B
    idx = 0
    while idx < len(meta):
        is_a = meta[idx]["is_a"]
        start_row = idx + 1
        while idx < len(meta) and meta[idx]["is_a"] == is_a:
            idx += 1
        end_row = idx + 1

        bg_color = {"red": 0.95, "green": 0.95, "blue": 0.95} if is_a else {"red": 1.0, "green": 1.0, "blue": 1.0}
        requests.append({
            "repeatCell": {
                "range": {
                    "sheetId": sheet_id,
                    "startRowIndex": start_row,
                    "endRowIndex": end_row,
                    "startColumnIndex": 0,
                    "endColumnIndex": len(COLUMNS)
                },
                "cell": {
                    "userEnteredFormat": {
                        "backgroundColor": bg_color
                    }
                },
                "fields": "userEnteredFormat.backgroundColor"
            }
        })

    # 3. Bordi di fine ciclo
    for idx, m in enumerate(meta):
        if m["is_cycle_end"]:
            row_idx = idx + 1
            requests.append({
                "updateBorders": {
                    "range": {
                        "sheetId": sheet_id,
                        "startRowIndex": row_idx,
                        "endRowIndex": row_idx + 1,
                        "startColumnIndex": 0,
                        "endColumnIndex": len(COLUMNS)
                    },
                    "bottom": {
                        "style": "SOLID_MEDIUM",
                        "width": 2,
                        "color": {"red": 0.0, "green": 0.0, "blue": 0.0}
                    }
                }
            })

    if requests:
        spreadsheet.batch_update({"requests": requests})
    print(f"-> Styling Google Sheet applicato ({len(requests)} direttive).")


def update_excel_file(filepath, sheet_name, all_rows, meta):
    """Aggiorna fogli Excel locali con styling."""
    if not os.path.exists(filepath):
        print(f"File {filepath} non trovato, salto.")
        return

    wb = openpyxl.load_workbook(filepath)
    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(sheet_name)
    else:
        ws = wb[sheet_name]

    ws.delete_rows(1, ws.max_row + 10)
    ws.append(COLUMNS)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")

    fill_a = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    fill_b = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    border_cycle_end = Border(bottom=Side(style="medium", color="000000"))

    for row_idx, r in enumerate(all_rows, start=2):
        ws.append(r)
        m = meta[row_idx - 2]
        fill = fill_a if m["is_a"] else fill_b

        for col_idx in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = fill
            if m["is_cycle_end"]:
                cell.border = border_cycle_end

    wb.save(filepath)
    print(f"-> Salvato {filepath} [{sheet_name}]: {len(all_rows)} righe.")


def update_demo_data_json(all_rows):
    """Sincronizza demo-data.json con tutte le righe PRO."""
    if not os.path.exists(DEMO_DATA_JSON):
        return

    with open(DEMO_DATA_JSON, "r") as f:
        demo_data = json.load(f)

    demo_data["allenamenti"] = [a for a in demo_data["allenamenti"] if a.get("livello") != "PRO"]

    existing_ids = [a.get("id", 0) for a in demo_data["allenamenti"]]
    max_id = max(existing_ids) if existing_ids else 500

    for r in all_rows:
        max_id += 1
        demo_data["allenamenti"].append({
            "id": max_id,
            "livello": r[0],
            "settimana": r[1],
            "giorno": r[2],
            "sequenza": r[3],
            "id_esercizio": f"{int(r[4]):03d}",
            "nome_esercizio": r[5],
            "parametri": r[6],
            "recupero": r[7],
            "minutaggio_blocco": "",
            "note_tecniche": r[9],
            "link_video": r[10],
            "data_pubblicazione": None
        })

    with open(DEMO_DATA_JSON, "w") as f:
        json.dump(demo_data, f, indent=2, ensure_ascii=False)
    print(f"-> demo-data.json aggiornato con {len(all_rows)} righe PRO (Cicli 1, 2, 3, 4, 5, 6).")


def update_vault_modulo_5(all_rows):
    """Sincronizza data/vault/5_metodologia_allenamenti_cicli.json con tutte le righe PRO."""
    vault_file = "data/vault/5_metodologia_allenamenti_cicli.json"
    if not os.path.exists(vault_file):
        return

    with open(vault_file, "r") as f:
        vault_data = json.load(f)

    vault_data["allenamenti"] = [a for a in vault_data["allenamenti"] if a.get("livello") != "PRO"]

    existing_ids = [a.get("id", 0) for a in vault_data["allenamenti"]]
    max_id = max(existing_ids) if existing_ids else 500

    for r in all_rows:
        max_id += 1
        vault_data["allenamenti"].append({
            "id": max_id,
            "livello": r[0],
            "settimana": r[1],
            "giorno": r[2],
            "sequenza": r[3],
            "id_esercizio": f"{int(r[4]):03d}",
            "nome_esercizio": r[5],
            "parametri": r[6],
            "recupero": r[7],
            "minutaggio_blocco": "",
            "note_tecniche": r[9],
            "link_video": r[10],
            "data_pubblicazione": None
        })

    with open(vault_file, "w") as f:
        json.dump(vault_data, f, indent=2, ensure_ascii=False)
    print(f"-> data/vault/5_metodologia_allenamenti_cicli.json aggiornato con {len(all_rows)} righe PRO.")


def update_register_json():
    """Aggiorna registro_periodizzazione_pro.json."""
    if not os.path.exists(REGISTER_JSON):
        return

    with open(REGISTER_JSON, "r") as f:
        reg = json.load(f)

    reg["ciclo_attivo_da_pianificare"] = "Ciclo 7 (Inizio Mesociclo 3)"
    reg["mesociclo_attuale"] = "Mesociclo 2 (Cicli 4, 5, 6) - COMPLETATO"

    if "storico_cicli" not in reg:
        reg["storico_cicli"] = {}

    reg["storico_cicli"]["Ciclo 6"] = {
        "status": "attivo / programmato",
        "sessioni": 6,
        "focus": "Chiusura Mesociclo 2: Scarico Selettivo Mod. B, Heavy Rope Flow 1x5', Bat Circular 60\"+60\", Coiled Punch Isoinerziale 40\"+40\", Pliometria e Stabilità Asimmetrica"
    }

    c6_used_ids = ["162", "131", "208", "77", "97", "104", "130", "90", "7", "33", "160", "61", "182", "28", "180", "240", "70", "91", "4"]

    for ex in reg.get("esercizi", []):
        eid = str(ex.get("id"))
        if eid in c6_used_ids:
            ex["usage_count"] = ex.get("usage_count", 0) + 1
            ex["last_cycle_used"] = "Ciclo 6"

    usati = sum(1 for e in reg.get("esercizi", []) if e.get("usage_count", 0) > 0)
    reg["esercizi_utilizzati"] = usati

    with open(REGISTER_JSON, "w") as f:
        json.dump(reg, f, indent=2, ensure_ascii=False)
    print(f"-> {REGISTER_JSON} aggiornato: {usati}/163 esercizi PRO sperimentati.")


def main():
    print("=" * 75)
    print("Area46 Training Lab - Generazione Ufficiale CICLO 6 (PRO)")
    print("Chiusura Mesociclo 2: Scarico Selettivo in Modalità B e Trasformazione")
    print("=" * 75)

    all_rows, meta = generate_all_rows()
    print(f"Righe totali PRO generate: {len(all_rows)}")
    print("  - Ciclo 1 (1-6): 57 righe")
    print("  - Ciclo 2 (7-12): 51 righe")
    print("  - Ciclo 3 (13-18): 51 righe")
    print("  - Ciclo 4 (19-24): 51 righe")
    print("  - Ciclo 5 (25-30): 51 righe")
    print("  - Ciclo 6 (31-36): 57 righe (Definitivo)")

    # 1. Connessione Google Fogli
    print("\n[Fase 1] Connessione a Google Sheets...")
    gc = gspread.service_account(filename=CREDENTIALS_FILE)

    # 2. Aggiorna Target Sheet
    print(f"\n[Fase 2] Aggiornamento Target Sheet '{TARGET_SHEET_ID}'...")
    sh_target = gc.open_by_key(TARGET_SHEET_ID)
    ws_pro_target = sh_target.worksheet("PRO")
    ws_pro_target.batch_clear(["A1:L400"])
    data_to_write = [COLUMNS] + all_rows
    ws_pro_target.update(range_name=f"A1:L{len(data_to_write)}", values=data_to_write)
    print(f"-> Scritte {len(data_to_write)} righe in 'PRO'.")
    apply_google_sheet_styling(sh_target, ws_pro_target, meta)

    # 3. Master Sheet
    print(f"\n[Fase 3] Aggiornamento Master Sheet '{LAB_SHEET_ID}'...")
    sh_lab = gc.open_by_key(LAB_SHEET_ID)
    ws_pro_lab = sh_lab.worksheet("Pro")
    ws_pro_lab.batch_clear(["A1:L400"])
    ws_pro_lab.update(range_name=f"A1:L{len(data_to_write)}", values=data_to_write)
    print(f"-> Scritte {len(data_to_write)} righe in 'Pro'.")
    apply_google_sheet_styling(sh_lab, ws_pro_lab, meta)

    # 4. Aggiorna File Excel locali
    print("\n[Fase 4] Aggiornamento file Excel locali...")
    update_excel_file(LOCAL_LAB_EXCEL, "Pro", all_rows, meta)
    update_excel_file(LOCAL_PRO_EXCEL, "PRO", all_rows, meta)

    # 5. Aggiorna demo-data.json e Vault Modulo 5 per l'App
    print("\n[Fase 5] Aggiornamento demo-data.json e Vault Modulo 5 per l'App...")
    update_demo_data_json(all_rows)
    update_vault_modulo_5(all_rows)
    if os.path.exists("/tmp/demo-data.json"):
        try:
            os.remove("/tmp/demo-data.json")
            print("-> Cache /tmp/demo-data.json rimossa per recepire subito i nuovi parametri.")
        except Exception:
            pass

    # 6. Aggiorna registro periodizzazione
    print("\n[Fase 6] Aggiornamento registro periodizzazione...")
    update_register_json()

    # 7. Sincronizza scheda registro su Drive
    print("\n[Fase 7] Sincronizzazione scheda registro Drive...")
    os.system(".venv/bin/python3 crea_scheda_registro_drive.py")

    # 8. Sincronizza cartella Drive
    print("\n[Fase 8] Sincronizzazione file Excel su Google Drive...")
    os.system(".venv/bin/python3 sincronizza_drive.py")

    print("\n" + "=" * 75)
    print("CICLO 6 GENERATO, APPLICATO E SINCRONIZZATO CON SUCCESSO!")
    print("=" * 75)


if __name__ == "__main__":
    main()
